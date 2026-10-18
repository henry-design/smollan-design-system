"""
Analyse every template slide in smollan_master.pptx and infer its fillable SLOTS.

    python tools/analyze_slots.py            -> tools/auto_specs.json

For each template slide it records every placeholder / text shape / table /
chart / picture with geometry, font size, the template's own instruction text
and an inferred ROLE (heading, subheading, intro, item heading, value, text,
note, photo, chart, table...). Repeated blocks (same role + same size) are
grouped into ITEMS in reading order. tools/gen_registry.py merges this with the
hand-written curation (tools/curation.py) to produce the RECIPES in smollan/recipes.py.
"""
import hashlib
import json
import os
import re
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from smollan import MASTER_PATH  # noqa: E402

E = 914400
# "CLIENT LOGO GOES HERE - DELETE & ADD CLIENT LOGO" placeholder images used in the packs
LOGO_IMAGES = {"2f5b81931b", "e38f73d330"}
WORDNUM = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7,
           "eight": 8, "nine": 9, "ten": 10}
DECOR = re.compile(r"^(\W*#\W*|WINNING AT THE POINT OF PURCHASE.*|0?\d{1,2}\.?|SMOLLAN|)$", re.S)


def inch(v):
    return round((v or 0) / E, 2)


def text_of(shape):
    if not shape.has_text_frame:
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()


def font_size(el, fallback):
    for tag in ("a:rPr", "a:defRPr", "a:endParaRPr"):
        for n in el.iter(qn(tag)):
            if n.get("sz"):
                return int(n.get("sz")) / 100
    return fallback


def is_caps(el, text):
    for n in el.iter(qn("a:rPr"), qn("a:defRPr")):
        if n.get("cap") == "all":
            return True
    letters = [c for c in text if c.isalpha()]
    return bool(letters) and sum(c.isupper() for c in letters) / len(letters) > 0.8


def line_limit(prompt):
    p = prompt.lower()
    m = re.search(r"(?:no )?(?:longer|more) than (\w+) lines?", p) or \
        re.search(r"max(?:imum)? (\w+) lines?", p) or \
        re.search(r"over (\w+) lines? only", p) or re.search(r"(\w+)[- ]line (?:heading|kpi|statement)", p) or \
        re.search(r"over (?:\w+ or )?(\w+) lines", p) or re.search(r"not exceed(?:ing)? (\w+) lines", p) or \
        re.search(r"(\w+) lines? of (?:copy|text)", p) or re.search(r"(\w+) lines? maximum", p)
    if not m:
        return None
    w = m.group(1)
    return WORDNUM.get(w, int(w) if w.isdigit() else None)


ROLE_RULES = [  # (role, regex on prompt/example text)
    ("presenter", r"presenter name"),
    ("date", r"edit date|xx september|\b20\d\d\b$"),
    ("subheading", r"sub-?heading|sub-head"),
    ("intro", r"introduct|intro text|insert intro"),
    ("value", r"^[\$R€£]?\s?[\d.,]+\s?(%|mil|m|k|bn)?$|^xxxx?$"),
    ("note", r"secondary text|short descriptor|footnote|note:"),
    ("label", r"statistic title|short heading|your short|highlight callout|^heading \w+$|^heading$|measure one|"
              r"title for your|^name surname|^role|^key|^step \d"),
    ("item_heading", r"heading|title|headline|concept|feature|differentiator|kpi"),
    ("bullets", r"bullet|master text styles|second level"),
    ("text", r"descript|paragraph|copy|text|detail|explain"),
]


def role_for(prompt, example, ph_type, is_title, y, w):
    if is_title:
        return "heading"
    if ph_type == "PICTURE":
        return "photo"
    if ph_type == "CHART":
        return "chart"
    if ph_type == "TABLE":
        return "table"
    if ph_type == "MEDIA_CLIP":
        return "video"
    if ph_type == "ORG_CHART":
        return "org_area"
    probe = (prompt or example or "").strip().lower()
    for role, rx in ROLE_RULES:
        if re.search(rx, probe):
            return role
    return "text"


def describe_shape(sh, lay_prompts, slide):
    d = {"id": sh.shape_id, "name": sh.name, "x": inch(sh.left), "y": inch(sh.top),
         "w": inch(sh.width), "h": inch(sh.height)}
    ex = text_of(sh)
    if sh.is_placeholder:
        pf = sh.placeholder_format
        d["ph"] = pf.idx
        d["ph_type"] = str(pf.type).split(".")[-1].split(" ")[0]
        lp = lay_prompts.get(pf.idx, {})
        d["prompt"] = lp.get("text", "")
        size = font_size(sh._element, None) or lp.get("size") or (28 if pf.idx == 0 else 12)
        if not d["w"]:
            d.update({k: lp.get(k, 0) for k in ("x", "y", "w", "h")})
    else:
        d["prompt"] = ""
        size = font_size(sh._element, 12)
    if getattr(sh, "has_table", False) and sh.has_table:
        tb = sh.table
        d["kind"] = "table"
        d["rows"], d["cols"] = len(tb.rows), len(tb.columns)
        d["merged"] = any(c.is_merge_origin or c.is_spanned for r in tb.rows for c in r.cells)
        d["cells"] = [[c.text.strip()[:40] for c in r.cells] for r in tb.rows]
        return d
    if getattr(sh, "has_chart", False) and sh.has_chart:
        ch = sh.chart
        d["kind"] = "chart"
        d["chart_type"] = str(ch.chart_type).split(" ")[0].split(".")[-1]
        try:
            plot = ch.plots[0]
            d["categories"] = [str(c) for c in plot.categories]
            d["series"] = [s.name for p in ch.plots for s in p.series]
        except Exception:
            d["categories"], d["series"] = [], []
        return d
    d["kind"] = "ph" if sh.is_placeholder else ("group" if sh.shape_type == 6 else "shape")
    d["example"] = ex[:200]
    d["size"] = size
    d["caps"] = is_caps(sh._element, d["prompt"] or ex)
    d["max_lines"] = line_limit(d["prompt"] or ex)
    return d


def layout_prompts(layout):
    out = {}
    for ph in layout.placeholders:
        out[ph.placeholder_format.idx] = {
            "text": text_of(ph)[:200], "size": font_size(ph._element, None),
            "x": inch(ph.left), "y": inch(ph.top), "w": inch(ph.width), "h": inch(ph.height)}
    return out


def group_items(elements):
    """Group repeated same-role elements into items (reading order)."""
    from collections import defaultdict
    buckets = defaultdict(list)
    for e in elements:
        if e["role"] in ("heading", "presenter", "date", "table", "chart"):
            continue
        key = (e["role"], round(e["w"] * 4) / 4, round(e["h"] * 4) / 4)
        buckets[key].append(e)
    reps = {k: v for k, v in buckets.items() if len(v) >= 2}
    if not reps:
        return None
    counts = defaultdict(int)
    for v in reps.values():
        counts[len(v)] += len(v)
    n = max(counts, key=lambda c: (counts[c], c))

    def reading(e):
        return (round(e["y"] / 0.6), e["x"])

    fields = {}
    for (role, _, _), v in sorted(reps.items(), key=lambda kv: min(e["y"] for e in kv[1])):
        if len(v) != n:
            continue
        name = role if role not in fields else f"{role}2"
        fields[name] = [e["ref"] for e in sorted(v, key=reading)]
    return {"n": n, "fields": fields}


def main():
    prs = Presentation(MASTER_PATH)
    out = {}
    for s in prs.slides:
        key = s._element.cSld.get("name")
        lp = layout_prompts(s.slide_layout)
        els = []
        for sh in s.shapes:
            if sh.shape_type == 13 and not sh.is_placeholder and                     hashlib.sha1(sh.image.blob).hexdigest()[:10] in LOGO_IMAGES:
                els.append({"ref": f"id:{sh.shape_id}", "id": sh.shape_id, "kind": "shape", "role": "logo",
                            "x": inch(sh.left), "y": inch(sh.top), "w": inch(sh.width), "h": inch(sh.height),
                            "prompt": "client logo (optional)"})
                continue
            d = describe_shape(sh, lp, s)
            if d["kind"] in ("shape", "group") and (not d.get("example") or DECOR.match(d["example"])):
                continue
            if d["kind"] == "group":
                continue
            d["ref"] = f"ph:{d['ph']}" if "ph" in d and d["ph"] != 4294967295 else f"id:{d['id']}"
            if d["kind"] in ("table", "chart"):
                d["role"] = d["kind"]
            else:
                d["role"] = role_for(d["prompt"], d.get("example", ""), d.get("ph_type"),
                                     d.get("ph") == 0 or d.get("ph_type") in ("TITLE", "CENTER_TITLE"),
                                     d["y"], d["w"])
            els.append(d)
        # pictures from the layout that the slide did not instantiate
        have = {e.get("ph") for e in els}
        for ph in s.slide_layout.placeholders:
            pf = ph.placeholder_format
            if pf.idx not in have and str(pf.type).split(".")[-1].startswith(("PICTURE", "CHART", "TABLE", "ORG", "TITLE", "CENTER_TITLE")):
                els.append({"ref": f"ph:{pf.idx}", "ph": pf.idx, "kind": "ph",
                            "ph_type": str(pf.type).split(".")[-1].split(" ")[0], "x": inch(ph.left),
                            "y": inch(ph.top), "w": inch(ph.width), "h": inch(ph.height),
                            "role": role_for("", "", str(pf.type).split(".")[-1].split(" ")[0],
                                             str(pf.type).split(".")[-1].startswith(("TITLE", "CENTER_TITLE")), 0, 0),
                            "prompt": text_of(ph)[:200], "size": font_size(ph._element, 24), "max_lines": line_limit(text_of(ph)),
                            "layout_only": True})
        out[key] = {"layout": s.slide_layout.name, "elements": els, "items": group_items(els)}
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "auto_specs.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=1, ensure_ascii=False)
    print("wrote", path, len(out))


if __name__ == "__main__":
    main()

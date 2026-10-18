"""
Remove client-specific content from smollan_master.pptx (run after build_master.py).

    python tools/scrub_master.py

The original packs contain worked examples from real client work (retailer names, hours,
costs, client programmes). This step makes the master safe to share publicly:
  * client / retailer / programme names -> neutral names (CLIENT, RETAILER A, ...)
  * figures in the Growth Teams / regional example tables -> zeros (format kept: 4 900 -> 0 000)
  * the full-photo differentiators slide -> the generic differentiator wording
  * document metadata (author, last modified by) -> neutral
It also scrubs tools/build_info.json (designer notes copied from the packs).
"""
import copy
import json
import os
import re
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from smollan import MASTER_PATH  # noqa: E402

# Client / retailer / programme names are NOT committed (the list itself would reveal client
# relationships). Put [[regex, replacement], ...] in tools/scrub_names.local.json (git-ignored).
_LOCAL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "scrub_names.local.json")
NAMES = json.load(open(_LOCAL, encoding="utf-8")) if os.path.exists(_LOCAL) else []
NUMBER_SLIDES = re.compile(r"^(p04gt\.s\d+|p04\.s04|p04\.s06)$")
NUM = re.compile(r"(R\s?)?\d[\d\s,.:]*\d%?|\d+%|\d+\*")


def scrub_names(text):
    for rx, rep in NAMES:
        def sub(m, rep=rep):
            return rep.upper() if m.group(0).isupper() else rep
        text = re.sub(rx, sub, text, flags=re.I)
    return text


def zero_numbers(text):
    return NUM.sub(lambda m: re.sub(r"\d", "0", m.group(0)), text)


def main():
    prs = Presentation(MASTER_PATH)
    slides = {s._element.cSld.get("name"): s for s in prs.slides}
    changed = 0
    parts = [s for s in prs.slides] + list(prs.slide_masters[0].slide_layouts)
    for obj in parts:
        key = obj._element.cSld.get("name") or ""
        for t in obj._element.iter(qn("a:t")):
            new = scrub_names(t.text or "")
            if NUMBER_SLIDES.match(key):
                new = zero_numbers(new)
            if new != t.text:
                t.text = new
                changed += 1
    # full-photo differentiators: copy the generic table text from the pale-panel variant
    src, dst = slides.get("p14.s04"), slides.get("p14.s06")
    if src is not None and dst is not None:
        src_tables = {sh.shape_id: sh for sh in src.shapes if getattr(sh, "has_table", False) and sh.has_table}
        for sh in dst.shapes:
            if getattr(sh, "has_table", False) and sh.has_table and sh.shape_id in src_tables:
                for r, row in enumerate(sh.table.rows):
                    for c, cell in enumerate(row.cells):
                        s_cell = src_tables[sh.shape_id].table.cell(r, c)
                        cell._tc.remove(cell._tc.get_or_add_txBody())
                        cell._tc.insert(0, copy.deepcopy(s_cell._tc.get_or_add_txBody()))
                        changed += 1
    cp = prs.core_properties
    cp.author = cp.last_modified_by = "Smollan Design System"
    cp.title, cp.comments, cp.keywords, cp.subject = "Smollan master template", "", "", ""
    prs.save(MASTER_PATH)
    info_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "build_info.json")
    if os.path.exists(info_path):
        info = json.load(open(info_path, encoding="utf-8"))
        for v in info.values():
            v["notes"] = [zero_numbers(scrub_names(n)) for n in v.get("notes", [])]
        json.dump(info, open(info_path, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    print(f"scrubbed {changed} text runs/cells in {MASTER_PATH}")


if __name__ == "__main__":
    main()

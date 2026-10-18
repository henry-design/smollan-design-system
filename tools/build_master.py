"""
Build smollan_master.pptx from the original Smollan PPT packs.

    python tools/build_master.py "G:/Other computers/My Computer/SMOLLAN/PPT Packs"

What it does (the originals are only read, never written):
  1. Starts from Smollan_Core-Template_2024.pptx (its master is identical to every
     pack's master apart from the layout list - verified).
  2. Copies every layout that an example slide uses (plus the hidden variants
     listed in HIDDEN_LAYOUTS) into that one master, de-duplicated by content.
  3. Copies every example slide (except each pack's "HOW TO USE THIS PACK"
     slide) as a template slide named "<pack>.s<NN>".
  4. "Promotes" content baked into layouts (tables, charts, fixed wording such
     as the Why-Smollan pillars) onto the template slides so it is editable.
  5. Removes designer sticky-notes ("EDIT TABLES IN MASTER", "Option 1", ...)
     and records them in build_info.json as usage notes.
"""
import hashlib
import json
import os
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from smollan._opc import (PartCopier, _add_slide_part, _dedupe_partnames, _renumber_shape_ids,  # noqa: E402
                          _strip_unused_rels)
from smollan.config import MASTER_PATH  # noqa: E402

PACKS = [  # (code, file name, slides to skip besides slide 1)
    ("core", "Smollan_Core-Template_2024.pptx", []),
    ("p01", "1  -  Smollan_Photo-Covers.pptx", []),
    ("p02", "2  -  Smollan_Photo+Client.pptx", [2]),
    ("p03", "3  -  Thank-You.pptx", []),
    ("p04", "4  -  Tables.pptx", []),
    ("p04gt", "4  -  Tables - GROWTH TEAMS.pptx", []),
    ("p05", "5  -  Graphs.pptx", [2]),
    ("p06", "6  -  Info-Buckets.pptx", []),
    ("p07", "7  -  Org-Charts.pptx", []),
    ("p08", "8  -  Client-Feedback.pptx", []),
    ("p09", "9  -  KPIs.pptx", []),
    ("p10", "10  -  Process.pptx", []),
    ("p11", "11  -  Infographic.pptx", []),
    ("p12", "12  -  Devices.pptx", []),
    ("p13", "13  -  Stats.pptx", []),
    ("p14", "14  -  Why-Smollan.pptx", []),
]
# Layouts present in a pack's master but not used by any example slide.
HIDDEN_LAYOUTS = {
    "p03": None,   # None = every unused layout of the pack (Thank-you A/B/C variants)
    "p11": ["HEAD + INTRO + Loop Infographic"],
}
DECOR = re.compile(r"^(\W*#\W*|WINNING AT THE POINT OF PURCHASE.*|0?\d{1,2}\.?|SMOLLAN|)$", re.S)
STICKY = re.compile(
    r"(EDIT TABLES? IN MASTER|^For Renewal clients|^Generic\s*\W\s*show to client|"
    r"Figures last updated|^Option \d\s*$|^Slide can be used to zone)", re.I | re.S)
REL_ATTRS = [qn("r:embed"), qn("r:id"), qn("r:link"), qn("r:pict")]


def text_of(el):
    return "".join(t.text or "" for t in el.iter(qn("a:t"))).strip()


def layout_hash(layout):
    xml = etree.tostring(layout._element)
    xml = re.sub(rb'r:(embed|id|link)="rId\d+"', b"", xml)
    h = hashlib.md5(xml)
    for rId, rel in sorted(layout.part.rels.items()):
        if not rel.is_external and rel.reltype != RT.SLIDE_MASTER:
            h.update(hashlib.md5(rel.target_part.blob).digest())
    return h.hexdigest()


def promotable(sp):
    """Layout shapes that carry content (not decoration) -> move to slide."""
    tag = etree.QName(sp).localname
    if sp.find(".//" + qn("p:ph")) is not None and tag != "grpSp":
        return False
    if tag == "graphicFrame":
        uri = sp.find(".//" + qn("a:graphicData")).get("uri", "")
        return uri.endswith("/table") or uri.endswith("/chart")
    txt = text_of(sp)
    if tag == "grpSp":
        texts = [text_of(s) for s in sp.iter(qn("p:sp"))]
        return any(t and not DECOR.match(t) for t in texts) or \
            any(g.find(".//" + qn("a:graphicData")) is not None for g in sp.iter(qn("p:graphicFrame")))
    return bool(txt) and not DECOR.match(txt)


def main(src_dir, only=None, promote=True, out=None):
    dest = Presentation(os.path.join(src_dir, PACKS[0][1]))
    master = dest.slide_masters[0]
    copier = PartCopier(dest.part.package)
    keep_alive, info = [], {}
    by_hash = {}   # layout hash -> dest layout part
    lid_lst = master._element.get_or_add_sldLayoutIdLst()

    def next_layout_id():
        ids = [int(e.get("id")) for e in lid_lst if e.get("id")] + [int(e.get("id")) for e in dest.part._element.iter(qn("p:sldMasterId"))]
        return max(ids) + 1

    # Core layouts: rename + register hashes.
    for lay in master.slide_layouts:
        by_hash[layout_hash(lay)] = lay.part
        lay._element.cSld.set("name", f"CORE | {lay.name}")

    # Core slides become template slides in place.
    for i, s in enumerate(dest.slides, 1):
        key = f"core.s{i:02d}"
        s._element.cSld.set("name", key)
        info[key] = {"pack": PACKS[0][1], "slide": i, "layout": s.slide_layout.name.split(" | ", 1)[1],
                     "hidden": False, "notes": []}
        for rId, rel in list(s.part.rels.items()):
            if rel.reltype == RT.NOTES_SLIDE:
                s.part.drop_rel(rId)

    for code, fname, skip in PACKS[1:]:
        if only and code not in only:
            continue
        src = Presentation(os.path.join(src_dir, fname))
        keep_alive.append(src)
        src_master_part = src.slide_masters[0].part
        redirect = {id(src_master_part): master.part}
        slides = [(i, s) for i, s in enumerate(src.slides, 1) if i != 1 and i not in skip]
        used_layouts = {id(s.slide_layout): s.slide_layout for _, s in slides}
        hidden = []
        if code in HIDDEN_LAYOUTS:
            want = HIDDEN_LAYOUTS[code]
            all_used = {id(s.slide_layout) for s in src.slides}
            for lay in src.slide_masters[0].slide_layouts:
                if id(lay) not in all_used and (want is None or lay.name in want):
                    hidden.append(lay)
        for lay in list(used_layouts.values()) + hidden:
            h = layout_hash(lay)
            if h in by_hash:
                redirect[id(lay.part)] = by_hash[h]
                continue
            new = copier.copy(lay.part, redirect)
            new._element.cSld.set("name", f"{code.upper()} | {lay.name}")
            rId = master.part.relate_to(new, RT.SLIDE_LAYOUT)
            nid = next_layout_id()
            e = etree.SubElement(lid_lst, qn("p:sldLayoutId"))
            e.set("id", str(nid))
            e.set(qn("r:id"), rId)
            by_hash[h] = new
            redirect[id(lay.part)] = new

        for i, s in slides:
            new = copier.copy(s.part, redirect)
            new._element.cSld.set("name", f"{code}.s{i:02d}")
            _add_slide_part(dest, new)
            info[f"{code}.s{i:02d}"] = {"pack": fname, "slide": i, "layout": s.slide_layout.name,
                                        "hidden": False, "notes": []}
        for lay in hidden:
            dl = redirect[id(lay.part)]
            layout_obj = [l for l in master.slide_layouts if l.part is dl][0]
            sl = dest.slides.add_slide(layout_obj)
            idx = list(src.slide_masters[0].slide_layouts).index(lay)
            key = f"{code}.h{idx:02d}"
            sl._element.cSld.set("name", key)
            info[key] = {"pack": fname, "slide": None, "layout": lay.name, "hidden": True, "notes": []}

    # ---- promotion + sticky-note cleanup ---------------------------------
    for s in (dest.slides if promote else []):
        key = s._element.cSld.get("name")
        lay = s.slide_layout
        lay_tree = lay._element.cSld.spTree
        slide_tree = s._element.cSld.spTree
        moved = []
        for sp in list(lay_tree)[2:]:
            if promotable(sp):
                moved.append(sp)
        if moved:
            insert_at = 2
            for sp in moved:
                el = sp if lay_tree is None else sp
                clone = etree.fromstring(etree.tostring(el))
                for node in clone.iter():
                    for attr in REL_ATTRS:
                        rid = node.get(attr)
                        if rid and rid in lay.part.rels:
                            rel = lay.part.rels[rid]
                            target = rel.target_part
                            if rel.reltype == RT.CHART and getattr(target, "_smollan_owned", False):
                                target = copier.copy(target)
                            target._smollan_owned = True
                            node.set(attr, s.part.relate_to(target, rel.reltype))
                slide_tree.insert(insert_at, clone)
                insert_at += 1
            info[key]["promoted"] = len(moved)
            info[key]["_lay"] = id(lay)
        for sp in list(slide_tree)[2:]:
            t = text_of(sp)
            if t and STICKY.search(t):
                info[key]["notes"].append(re.sub(r"\s+", " ", t))
                slide_tree.remove(sp)
        _renumber_shape_ids(slide_tree)
        _strip_unused_rels(s.part, s._element)

    # Remove promoted shapes from layouts only after all slides cloned them.
    done = set()
    for s in (dest.slides if promote else []):
        lay = s.slide_layout
        if id(lay) in done:
            continue
        done.add(id(lay))
        lay_tree = lay._element.cSld.spTree
        removed = False
        for sp in list(lay_tree)[2:]:
            if promotable(sp):
                lay_tree.remove(sp)
                removed = True
        if removed:
            _strip_unused_rels(lay.part, lay._element)

    for v in info.values():
        v.pop("_lay", None)
    _dedupe_partnames(dest.part.package)
    dest.save(out or MASTER_PATH)
    with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "build_info.json"), "w", encoding="utf-8") as f:
        json.dump(info, f, indent=1, ensure_ascii=False)
    print(f"saved {MASTER_PATH}: {len(dest.slides)} template slides, "
          f"{len(master.slide_layouts)} layouts")


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("src", nargs="?", default=r"G:\Other computers\My Computer\SMOLLAN\PPT Packs")
    ap.add_argument("--only", help="comma list of pack codes (debug)")
    ap.add_argument("--no-promote", action="store_true")
    ap.add_argument("--out")
    a = ap.parse_args()
    main(a.src, a.only.split(",") if a.only else None, not a.no_promote, a.out)

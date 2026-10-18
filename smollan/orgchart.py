"""Org charts drawn from tree data in Smollan styling."""
from __future__ import annotations

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .brand import Brand
from .errors import SmollanError
from .shapes import _crop_to_fill


ORG_LEVEL_COLORS = [Brand.DARK_BLUE, Brand.GREEN, Brand.BLUE, Brand.CHARCOAL, Brand.CYAN_BLUE]


def _org_nodes(tree, depth=0, out=None):
    out = [] if out is None else out
    out.append((depth, tree))
    for c in tree.get("children") or []:
        _org_nodes(c, depth + 1, out)
    return out


def org_size(tree):
    return len(_org_nodes(tree))


def draw_org_chart(slide, x, y, w, h, tree, max_cols=8):
    """Draw an org chart inside a box (EMU). tree = {"name", "role", "photo"?, "children": [...]}.

    A manager whose reports are all leaves gets them STACKED underneath when the chart
    would otherwise be wider than `max_cols` boxes. Raises SmollanError when it cannot
    fit - then split the tree over two slides (top levels on one, each branch on its own).
    """
    nodes = _org_nodes(tree)
    if max(d for d, _ in nodes) > 4:
        raise SmollanError("org chart deeper than 5 levels - split it over several slides "
                           "(top structure on one slide, each department on its own).")

    def leaves_only(n):
        ch = n.get("children") or []
        return bool(ch) and all(not c.get("children") for c in ch)

    stack = set()

    def units(n):
        ch = n.get("children") or []
        if not ch or id(n) in stack:
            return 1
        return sum(units(c) for c in ch)

    candidates = sorted([n for d, n in nodes if d >= 1 and leaves_only(n)], key=lambda n: -len(n["children"]))
    while units(tree) > max_cols and candidates:
        stack.add(id(candidates.pop(0)))
    total_units = units(tree)
    if total_units > max_cols:
        raise SmollanError(f"org chart needs {total_units} columns (max {max_cols}). Split it: top levels on "
                           f"one slide, each branch on its own slide (pass the branch as the tree).")
    gap_x = int(Inches(0.15))
    box_w = int(min(Inches(2.6), (w - gap_x * (total_units - 1)) / total_units))
    box_h, stack_h, stack_gap = int(Inches(0.7)), int(Inches(0.5)), int(Inches(0.08))

    row_levels = 0

    def row_depth(n, d):
        nonlocal row_levels
        row_levels = max(row_levels, d)
        if id(n) not in stack:
            for c in n.get("children") or []:
                row_depth(c, d + 1)

    row_depth(tree, 0)
    max_stack = max([len(n["children"]) for d, n in nodes if id(n) in stack] + [0])
    stack_block = (int(Inches(0.1)) + max_stack * (stack_h + stack_gap)) if max_stack else 0
    free = h - (row_levels + 1) * box_h - stack_block
    if free < row_levels * Inches(0.25):
        raise SmollanError("org chart is too tall for the slide - split it over two slides, or remove a level.")
    gap_y = int(min(Inches(0.6), free / max(row_levels, 1))) if row_levels else 0
    used_w = total_units * box_w + (total_units - 1) * gap_x
    used_h = (row_levels + 1) * box_h + row_levels * gap_y + stack_block
    x0, y0 = x + int((w - used_w) / 2), y + int(max(h - used_h, 0) / 2)
    shapes = slide.shapes
    pos = {}

    def place(n, d, left_unit):
        u = units(n)
        cx = x0 + int((left_unit + u / 2) * (box_w + gap_x) - gap_x / 2)
        top = y0 + d * (box_h + gap_y)
        pos[id(n)] = (cx, top, d)
        if id(n) in stack:
            for i, c in enumerate(n.get("children") or []):
                pos[id(c)] = (cx, top + box_h + int(Inches(0.1)) + i * (stack_h + stack_gap), d + 1)
            return
        lu = left_unit
        for c in n.get("children") or []:
            place(c, d + 1, lu)
            lu += units(c)

    place(tree, 0, 0)

    def line(x1, y1, x2, y2):
        ln = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        ln.line.color.rgb = Brand.GREY_BLUE
        ln.line.width = Pt(1)

    def box(n, cx, top, d, small=False):
        bw = box_w if not small else int(box_w * 0.9)
        bh = box_h if not small else stack_h
        left = int(cx - bw / 2)
        photo = n.get("photo")
        if photo:
            pic = shapes.add_picture(str(photo), left, top, bh, bh)
            _crop_to_fill(pic, bh, bh)
            left, bw = left + bh, bw - bh
        sp = shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, bw, bh)
        sp.fill.solid()
        sp.fill.fore_color.rgb = ORG_LEVEL_COLORS[min(d, len(ORG_LEVEL_COLORS) - 1)]
        sp.line.fill.background()
        sp.shadow.inherit = False
        tf = sp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(n.get("name", "")).upper()
        r.font.size, r.font.bold = Pt(9 if not small else 8), True
        r.font.name, r.font.color.rgb = Brand.HEADING_FONT, Brand.WHITE
        if n.get("role"):
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            r2.text = str(n["role"])
            r2.font.size = Pt(8 if not small else 7)
            r2.font.name, r2.font.color.rgb = Brand.BODY_FONT, Brand.WHITE

    def draw(n):
        cx, top, d = pos[id(n)]
        kids = n.get("children") or []
        if kids and id(n) in stack:
            line(cx, top + box_h, cx, pos[id(kids[-1])][1] + int(Inches(0.05)))
            for c in kids:
                ccx, ctop, cd = pos[id(c)]
                box(c, ccx, ctop, cd, small=True)
        elif kids:
            mid = top + box_h + int(gap_y / 2)
            line(cx, top + box_h, cx, mid)
            xs = [pos[id(c)][0] for c in kids]
            if len(xs) > 1:
                line(min(xs), mid, max(xs), mid)
            for c in kids:
                line(pos[id(c)][0], mid, pos[id(c)][0], pos[id(c)][1])
                draw(c)
        box(n, cx, top, d)

    draw(tree)

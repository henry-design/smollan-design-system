"""Native, editable charts in Smollan colours (incl. an emulated waterfall)."""
from __future__ import annotations

import copy

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.oxml.ns import qn
from pptx.util import Pt

from .brand import Brand
from .errors import SmollanError
from .shapes import _remove_shape


CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_100": XL_CHART_TYPE.COLUMN_STACKED_100,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "stacked_bar_100": XL_CHART_TYPE.BAR_STACKED_100,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "line_plain": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
    "stacked_area": XL_CHART_TYPE.AREA_STACKED,
    "waterfall": XL_CHART_TYPE.COLUMN_STACKED,   # emulated with an invisible base series
}
_FAMILY = {  # python-pptx chart_type names -> our type names
    "COLUMN_CLUSTERED": "column", "COLUMN_STACKED_100": "stacked_100", "COLUMN_STACKED": "stacked_column",
    "BAR_CLUSTERED": "bar", "BAR_STACKED_100": "stacked_bar_100", "BAR_STACKED": "stacked_bar",
    "LINE_MARKERS_STACKED": "line", "LINE_MARKERS": "line", "LINE": "line_plain", "PIE": "pie",
    "DOUGHNUT": "doughnut", "AREA_STACKED": "stacked_area", "AREA": "area",
}


def _chart_family(chart):
    try:
        name = str(chart.chart_type).split(" ")[0].split(".")[-1]
    except Exception:
        return None
    return _FAMILY.get(name)


def _waterfall_data(values, totals):
    base, up, down, tot = [], [], [], []
    running = 0
    for i, v in enumerate(values):
        if i in totals:
            t = running if v is None else v
            running = t
            base.append(0); up.append(0); down.append(0); tot.append(t)
        elif v >= 0:
            base.append(running); up.append(v); down.append(0); tot.append(0)
            running += v
        else:
            running += v
            base.append(running); up.append(0); down.append(-v); tot.append(0)
    return [("Base", base), ("Increase", up), ("Decrease", down), ("Total", tot)]


def _chart_data(spec):
    cats = spec.get("categories")
    series = spec.get("series")
    if not cats or not series:
        raise SmollanError('chart needs "categories" and "series", e.g. '
                           '{"type": "column", "categories": ["Q1", "Q2"], "series": {"2026": [1, 2]}}')
    if isinstance(series, (list, tuple)):
        series = {spec.get("series_name", "Series 1"): series}
    data = CategoryChartData(number_format=spec.get("number_format", "General"))
    data.categories = cats
    if spec.get("type") == "waterfall":
        vals = list(series.values())[0]
        totals = set(spec.get("totals", [len(vals) - 1]))
        for name, vals2 in _waterfall_data(vals, totals):
            data.add_series(name, vals2)
        return data
    for name, vals in series.items():
        if len(vals) != len(cats):
            raise SmollanError(f"chart series {name!r} has {len(vals)} values for {len(cats)} categories")
        data.add_series(name, vals)
    return data


def style_chart(chart, spec):
    """Apply Smollan styling to a newly created chart."""
    ctype = spec.get("type", "column")
    chart.font.size = Pt(10)
    chart.font.name = Brand.BODY_FONT
    chart.font.color.rgb = Brand.CHARCOAL
    plot = chart.plots[0]
    multi = len(list(plot.series)) > 1 or ctype in ("pie", "doughnut")
    chart.has_legend = multi and ctype != "waterfall"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if ctype in ROUND_TYPES:
        pass                                    # slices, legend and labels: _style_round_chart below
    else:
        colors = Brand.SERIES
        if ctype == "waterfall":
            colors = [None, Brand.GREEN, Brand.BRIGHT_CYAN, Brand.BLUE]
            plot.gap_width = 60
            plot.overlap = 100
        for i, s in enumerate(plot.series):
            col = colors[i % len(colors)]
            fmt = s.format
            if col is None:
                fmt.fill.background()
                fmt.line.fill.background()
                continue
            if ctype.startswith("line"):
                fmt.line.color.rgb = col
                fmt.line.width = Pt(2.25)
                s.smooth = False
            else:
                fmt.fill.solid()
                fmt.fill.fore_color.rgb = col
        try:
            va = chart.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = Brand.PALE
            va.format.line.fill.background()
            chart.category_axis.format.line.color.rgb = Brand.GREY_BLUE
        except Exception:
            pass
    if ctype == "waterfall":
        fmt = spec.get("number_format", "#,##0")
        for i, s in enumerate(plot.series):
            if i == 0:
                continue                        # invisible base series: no labels
            dl = s.data_labels
            dl.show_value = True
            dl.font.size = Pt(9)
            dl.number_format = f"{fmt};-{fmt};;"  # hide zero labels
            dl.number_format_is_linked = False
            dl.font.color.rgb = Brand.WHITE if i == 3 else Brand.NEAR_BLACK
    elif ctype in ROUND_TYPES:
        _style_round_chart(chart, spec, new=True)
    elif spec.get("labels"):
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_value = True
        dl.font.size = Pt(9)
        dl.number_format = spec.get("number_format", "General")
        dl.number_format_is_linked = False
    if spec.get("title"):
        _set_chart_title(chart, spec["title"], Pt(12))
    else:
        chart.has_title = False


# ---- pie / doughnut ---------------------------------------------------------------------------
ROUND_TYPES = ("pie", "doughnut")
MAX_SLICES = 8


def _check_round(spec, family):
    """Guard-rails for pie/doughnut data (raised before the slide is built)."""
    if family not in ROUND_TYPES:
        return
    series = spec["series"]
    if isinstance(series, (list, tuple)):
        series = {"Series 1": series}
    n = len(spec["categories"])
    if n > MAX_SLICES:
        raise SmollanError(f"{family} charts hold max {MAX_SLICES} slices; you passed {n}. Group the smallest "
                           f"into 'Other', or use a bar chart ('type': 'bar').")
    if family == "pie" and len(series) > 1:
        raise SmollanError(f"a pie chart shows one series; you passed {len(series)}. Use 'type': 'doughnut' "
                           f"for concentric rings, two charts (e.g. graphs.2_intro), or a bar chart.")
    for name, vals in series.items():
        if any(v is not None and v < 0 for v in vals):
            raise SmollanError(f"{family} values cannot be negative (series {name!r}). Use a bar or waterfall chart.")
        if not any(v for v in vals if v is not None):
            raise SmollanError(f"{family} series {name!r} is all zero - nothing to draw.")


def _brand_slice_colours(chart, keep_existing):
    """Give every slice a Smollan colour and a white edge. Template slices keep their own colours;
    slices beyond the template's (it only styles its example slices) get the Brand.SERIES order."""
    for plot in chart.plots:
        for s in plot.series:
            n = len(list(s.values))
            ser = s._element
            for dpt in ser.findall(qn("c:dPt")):                  # drop formatting for points that no longer exist
                if int(dpt.find(qn("c:idx")).get("val")) >= n:
                    ser.remove(dpt)
            styled = {int(d.find(qn("c:idx")).get("val")) for d in ser.findall(qn("c:dPt"))}
            for i in range(n):
                if keep_existing and i in styled:
                    continue
                fmt = s.points[i].format
                fmt.fill.solid()
                fmt.fill.fore_color.rgb = Brand.SERIES[i % len(Brand.SERIES)]
                fmt.line.color.rgb = Brand.WHITE
                fmt.line.width = Pt(1.5)


def _normalise_legend(chart):
    """Templates format legend entries 1-4 individually and fix the legend box size, so a 5th+ category
    gets a fainter default style or is cut off. Use entry 1's style for all and let the legend auto-size."""
    leg = chart._chartSpace.find(".//" + qn("c:legend"))
    if leg is None:
        return
    entries = leg.findall(qn("c:legendEntry"))
    style = None
    for e in entries:
        if style is None and e.find(qn("c:txPr")) is not None:
            style = copy.deepcopy(e.find(qn("c:txPr")))
        leg.remove(e)
    layout = leg.find(qn("c:layout"))
    if layout is not None:
        leg.remove(layout)
    if style is not None:
        old = leg.find(qn("c:txPr"))
        if old is not None:
            leg.remove(old)
        ext = leg.find(qn("c:extLst"))
        if ext is not None:
            ext.addprevious(style)
        else:
            leg.append(style)


def _style_round_chart(chart, spec, new):
    plot = chart.plots[0]
    if new:
        _brand_slice_colours(chart, keep_existing=False)
        hole = chart._chartSpace.find(".//" + qn("c:holeSize"))
        if hole is not None:
            hole.set("val", "55")
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    else:
        _brand_slice_colours(chart, keep_existing=True)
        _normalise_legend(chart)
    labels = spec.get("labels", True if new else None)
    if labels is None:                                   # template chart: keep its own label setup
        return
    plot.has_data_labels = bool(labels)
    # Templates set labels per series, which overrides the plot level - so configure both.
    targets = [plot.data_labels] + [s.data_labels for s in plot.series] if labels else []
    for s in plot.series:
        ser_dl = s._element.find(qn("c:dLbls"))
        if not labels and ser_dl is not None:
            s._element.remove(ser_dl)
    for dl in targets:
        dl.show_value = labels != "percent"
        dl.show_percentage = labels == "percent"
        dl.show_category_name = False
        if labels == "percent" and not spec.get("number_format"):
            dl.number_format, dl.number_format_is_linked = "0%", False
        elif spec.get("number_format"):
            dl.number_format, dl.number_format_is_linked = spec["number_format"], False
        if new:
            dl.font.size, dl.font.bold = Pt(12), True
            dl.font.color.rgb = Brand.WHITE
            if chart.chart_type == XL_CHART_TYPE.PIE:
                dl.position = XL_LABEL_POSITION.INSIDE_END
    if new and chart.has_legend:
        chart.legend.font.size = Pt(11)


def _set_chart_title(chart, text, size=None):
    """Set a chart title, keeping the template's title style (python-pptx would reset it)."""
    chart.has_title = True
    title_el = chart.chart_title._title
    model = title_el.find(".//" + qn("c:txPr") + "//" + qn("a:defRPr"))
    if model is None:
        model = title_el.find(".//" + qn("a:rPr"))
    model = copy.deepcopy(model) if model is not None else None
    tf = chart.chart_title.text_frame
    tf.text = str(text)
    for p in tf.paragraphs:
        for r in p.runs:
            rpr = r._r.get_or_add_rPr()
            if model is not None:
                for k, v in model.attrib.items():
                    rpr.set(k, v)
                for child in model:
                    rpr.append(copy.deepcopy(child))
            elif size:
                r.font.size, r.font.bold = size, True


def put_chart(slide, shape, spec):
    """Fill a chart slot. Re-uses the template chart (and its styling) when the type matches."""
    if not isinstance(spec, dict):
        raise SmollanError("chart slot expects a dict spec")
    data = _chart_data(spec)
    want = spec.get("type")
    existing = shape.chart if getattr(shape, "has_chart", False) and shape.has_chart else None
    family = want or (_chart_family(existing) if existing is not None else "column")
    _check_round(spec, family)
    if existing is not None and (want is None or want == _chart_family(existing)):
        existing.replace_data(data)
        if family in ROUND_TYPES:
            _style_round_chart(existing, spec, new=False)
        elif "labels" in spec:
            plot = existing.plots[0]
            plot.has_data_labels = bool(spec["labels"])
            if spec["labels"]:
                plot.data_labels.show_value = True
        if spec.get("title"):
            _set_chart_title(existing, spec["title"])
        return shape
    xl = CHART_TYPES.get(want or "column")
    if xl is None:
        raise SmollanError(f"unknown chart type {want!r}; use one of {sorted(CHART_TYPES)}")
    if shape.is_placeholder and hasattr(shape, "insert_chart") and existing is None:
        frame = shape.insert_chart(xl, data)
    else:
        x, y, w, h = shape.left, shape.top, shape.width, shape.height
        _remove_shape(shape)
        frame = slide.shapes.add_chart(xl, x, y, w, h, data)
    style_chart(frame.chart, {**spec, "type": want or "column"})
    return frame

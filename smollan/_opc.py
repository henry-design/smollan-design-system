"""Low-level package helpers: copy parts between / within presentations."""
from __future__ import annotations

import hashlib
import re
from typing import Any

from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TARGET_MODE as RTM
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import PartFactory, _Relationship
from pptx.opc.packuri import PackURI
from pptx.parts.image import ImagePart


def _partname_template(partname: str) -> str:
    """'/ppt/charts/chart12.xml' -> '/ppt/charts/chart%d.xml'."""
    m = re.match(r"^(.*?)(\d*)(\.[A-Za-z0-9]+)$", partname)
    return f"{m.group(1)}%d{m.group(3).lower()}"


class PartCopier:
    """Deep-copies OPC parts (slides, layouts, charts, media...) into a package.

    redirect maps a source part to an existing destination part (e.g. a source
    layout -> its already-copied twin) so the graph is re-wired, not duplicated.
    """

    def __init__(self, dest_pkg):
        self.pkg = dest_pkg
        self.used = {str(p.partname) for p in dest_pkg.iter_parts()}
        self.memo: dict[int, Any] = {}
        self._images = None

    def next_partname(self, tmpl: str) -> PackURI:
        i = 1
        while (tmpl % i) in self.used:
            i += 1
        name = tmpl % i
        self.used.add(name)
        return PackURI(name)

    def copy(self, src, redirect: dict | None = None, skip_reltypes=(RT.NOTES_SLIDE,)):
        redirect = redirect or {}
        if id(src) in redirect:
            return redirect[id(src)]
        if id(src) in self.memo:
            return self.memo[id(src)]
        if isinstance(src, ImagePart) or str(src.partname).startswith("/ppt/media/"):
            if self._images is None:
                self._images = {hashlib.sha1(p.blob).hexdigest(): p for p in self.pkg.iter_parts()
                                if str(p.partname).startswith("/ppt/media/")}
            sha = hashlib.sha1(src.blob).hexdigest()
            if sha not in self._images:
                self._images[sha] = PartFactory(self.next_partname(_partname_template(str(src.partname))),
                                                src.content_type, self.pkg, src.blob)
            self.memo[id(src)] = self._images[sha]
            return self._images[sha]
        new = PartFactory(self.next_partname(_partname_template(str(src.partname))),
                          src.content_type, self.pkg, src.blob)
        self.memo[id(src)] = new
        for rId, rel in list(src.rels.items()):
            if rel.reltype in skip_reltypes:
                continue
            if rel.is_external:
                r = _Relationship(new.partname.baseURI, rId, rel.reltype, RTM.EXTERNAL, rel.target_ref)
            else:
                tgt = self.copy(rel.target_part, redirect, skip_reltypes)
                r = _Relationship(new.partname.baseURI, rId, rel.reltype, RTM.INTERNAL, tgt)
            new.rels._rels[rId] = r
        new.rels.__dict__.pop("_rels_by_reltype", None)
        return new


def _dedupe_partnames(pkg):
    """Give every part a unique name (python-pptx and PartCopier number independently)."""
    parts = list(pkg.iter_parts())
    used = {str(p.partname) for p in parts}
    seen = set()
    for p in parts:
        name = str(p.partname)
        if name in seen:
            tmpl, i = _partname_template(name), 1
            while (tmpl % i) in used:
                i += 1
            p.partname = PackURI(tmpl % i)
            used.add(tmpl % i)
        seen.add(str(p.partname))


def _strip_unused_rels(part, xml_el):
    """Drop relationships a slide no longer references (after removing shapes)."""
    xml = etree.tostring(xml_el).decode()
    for rId, rel in list(part.rels.items()):
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE, RT.SLIDE_MASTER):
            continue
        if f'"{rId}"' not in xml:
            part.drop_rel(rId)


def _renumber_shape_ids(sptree_el):
    """Make every cNvPr id in a shape tree unique (PowerPoint repairs duplicates)."""
    used = set()
    nodes = [n for n in sptree_el.iter() if n.tag.endswith("}cNvPr")]
    nxt = max([int(n.get("id", "0")) for n in nodes] + [1]) + 1
    for n in nodes:
        i = int(n.get("id", "0"))
        if i in used or i == 0:
            n.set("id", str(nxt))
            i = nxt
            nxt += 1
        used.add(i)


def _add_slide_part(prs, slide_part):
    rId = prs.part.relate_to(slide_part, RT.SLIDE)
    prs.slides._sldIdLst._add_sldId(id=prs.slides._sldIdLst._next_id, rId=rId)
    return rId


def _remove_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sldId in list(lst):
        if prs.part.related_part(sldId.rId) is slide.part:
            prs.part.drop_rel(sldId.rId)
            lst.remove(sldId)
            return

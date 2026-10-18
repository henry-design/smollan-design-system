"""Slot lookup (ph:/id:/cell: references), shape removal, photos and video."""
from __future__ import annotations

import os

from lxml import etree
from pptx.oxml.ns import qn

from .errors import SmollanError


def _iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # group
            yield from _iter_shapes(sh.shapes)


def _purge_alternate_content(slide, idx):
    """Remove mc:AlternateContent blocks (Office 2016 charts etc.) holding placeholder idx.

    python-pptx does not see shapes inside mc:AlternateContent, so they would otherwise
    survive underneath the content we add."""
    tree = slide._element.cSld.spTree
    for node in list(tree):
        if etree.QName(node).localname != "AlternateContent":
            continue
        for ph in node.iter(qn("p:ph")):
            if int(ph.get("idx", "0")) == idx:
                tree.remove(node)
                break


def _find_shape(slide, ref):
    """Resolve a slot reference: 'ph:<idx>' | 'id:<shape id>' | 'cell:<shape id>:<row>:<col>'."""
    kind, _, rest = ref.partition(":")
    if kind == "ph":
        idx = int(rest)
        for sh in slide.shapes:
            if sh.is_placeholder and sh.placeholder_format.idx == idx:
                return sh
        _purge_alternate_content(slide, idx)
        for lph in slide.slide_layout.placeholders:      # instantiate from the layout
            if lph.placeholder_format.idx == idx:
                slide.shapes.clone_placeholder(lph)
                for sh in slide.shapes:
                    if sh.is_placeholder and sh.placeholder_format.idx == idx:
                        return sh
        return None
    if kind == "id":
        sid = int(rest)
        for sh in _iter_shapes(slide.shapes):
            if sh.shape_id == sid:
                return sh
        return None
    if kind == "cell":
        sid, r, c = (int(x) for x in rest.split(":"))
        sh = _find_shape(slide, f"id:{sid}")
        return sh.table.cell(r, c) if sh is not None else None
    raise SmollanError(f"bad slot reference {ref!r}")


def _remove_shape(sh):
    """Remove a shape; if it sits inside mc:AlternateContent (e.g. Office 2016 charts), remove the whole block."""
    el = sh._element
    node = el
    while node.getparent() is not None and etree.QName(node.getparent()).localname in ("Choice", "Fallback", "AlternateContent"):
        node = node.getparent()
    node.getparent().remove(node)


def _crop_to_fill(pic, box_w, box_h):
    iw, ih = pic.image.size
    if not iw or not ih:
        return
    img_ratio, box_ratio = iw / ih, box_w / box_h
    if img_ratio > box_ratio:     # too wide -> crop left/right
        excess = 1 - box_ratio / img_ratio
        pic.crop_left = pic.crop_right = excess / 2
    else:
        excess = 1 - img_ratio / box_ratio
        pic.crop_top = pic.crop_bottom = excess / 2


def put_photo(slide, shape, path):
    """Fill a picture placeholder (or any shape's box) with an image, cropped to fill."""
    path = str(path)
    if not os.path.exists(path):
        raise SmollanError(f"photo not found: {path}")
    if shape.is_placeholder and hasattr(shape, "insert_picture"):
        return shape.insert_picture(path)
    x, y, w, h = shape.left, shape.top, shape.width, shape.height
    pic = slide.shapes.add_picture(path, x, y, w, h)
    _crop_to_fill(pic, w, h)
    _remove_shape(shape)
    return pic


def put_logo(slide, shape, path):
    """Place a logo inside a shape's box, scaled to fit (never cropped or stretched), centred."""
    path = str(path)
    if not os.path.exists(path):
        raise SmollanError(f"logo not found: {path}")
    x, y, w, h = shape.left, shape.top, shape.width, shape.height
    pic = slide.shapes.add_picture(path, x, y)
    iw, ih = pic.width, pic.height
    scale = min(w / iw, h / ih)
    pic.width, pic.height = int(iw * scale), int(ih * scale)
    pic.left, pic.top = int(x + (w - pic.width) / 2), int(y + (h - pic.height) / 2)
    _remove_shape(shape)
    return pic


def put_video(slide, shape, path, poster=None):
    x, y, w, h = shape.left, shape.top, shape.width, shape.height
    mov = slide.shapes.add_movie(str(path), x, y, w, h, poster_frame_image=poster, mime_type="video/mp4")
    _remove_shape(shape)
    return mov
